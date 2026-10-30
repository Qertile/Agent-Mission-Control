"""Task list for Linux: the counterpart of show-hud.ps1 and sync-sheet.ps1.

tasks.md parsing and ranking, the Status column of TaskDatabase.xlsx, edits made
from Agent Mission Control, and the Excel sync. Agent Mission Control imports
this module, so the window and the popup always agree on order and status.

    python3 tasks.py sync     rewrite TaskDatabase.xlsx from tasks.md
    python3 tasks.py popup    the standalone "Todo List" window
    python3 tasks.py list     print the list the popup shows
"""

import os
import re
import sys
from datetime import date, datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
TASKS_PATH = os.path.join(HERE, 'tasks.md')
XLSX_PATH = os.path.join(HERE, 'TaskDatabase.xlsx')

RESULT_LINE = re.compile(r'^\s+→\s*(?:已回覆|結果)')
RESULT_VALUE = re.compile(r'^\s+→\s*(?:已回覆|結果)[^:：]*[:：]\s*(.+)$')


def strip_tags(body):
    return re.sub(r'\[priority:\S+\]', '', re.sub(r'\[due:\S+\]', '', re.sub(r'\(added:\S+\)', '', body))).strip()


def read_lines(path):
    if not os.path.exists(path):
        return []
    with open(path, encoding='utf-8-sig') as f:
        return f.read().splitlines()


def unchecked_tasks(path=TASKS_PATH):
    """Open tasks with their description lines, as the popup shows them."""
    lines = read_lines(path)
    tasks, current = [], None
    i = 0
    while i < len(lines):
        line = lines[i]
        m = re.match(r'^- \[ \] (.+)$', line)
        if m:
            if current:
                tasks.append(current)
            body = m.group(1)
            due = re.search(r'\[due:(\S+)\]', body)
            added = re.search(r'\(added:(\S+)\)', body)
            current = {'text': strip_tags(body), 'description': [],
                       'due': due.group(1) if due else None,
                       'added': added.group(1) if added else None}
        elif re.match(r'^- \[x\]', line, re.I):
            if current:
                tasks.append(current)
            current = None
        elif re.match(r'^\s+\S', line) and current:
            if RESULT_LINE.match(line):
                # A result line by itself (no preceding description line owns it) - skip it.
                pass
            elif i + 1 < len(lines) and RESULT_LINE.match(lines[i + 1]):
                # A description line followed by a result line is a finished sub-item: skip both.
                i += 1
            else:
                current['description'].append(line.strip())
        else:
            if current:
                tasks.append(current)
            current = None
        i += 1
    if current:
        tasks.append(current)
    return tasks


def finished_tasks(path=TASKS_PATH):
    """The ticked tasks, latest finished first: {text, added, done, results: [{subitem,
    result}]}. done is the newest date on its result lines ('' when it has none; those last)."""
    out, current = [], None
    for line in read_lines(path):
        m = re.match(r'^- \[( |x)\] (.+)$', line, re.I)
        if m:
            current = None
            if m.group(1).lower() == 'x':
                added = re.search(r'\(added:(\S+)\)', m.group(2))
                current = {'text': strip_tags(m.group(2)), 'added': added.group(1) if added else '',
                           'done': '', 'results': []}
                out.append(current)
        elif current and re.match(r'^\s+\S', line):
            result = RESULT_VALUE.match(line)
            if result:
                day = re.search(r'\((\d{4}-\d{2}-\d{2})\)', line)
                current['done'] = max(current['done'], day.group(1) if day else '')
                if current['results'] and not current['results'][-1]['result']:
                    current['results'][-1]['result'] = result.group(1).strip()
                else:
                    current['results'].append({'subitem': '', 'result': result.group(1).strip()})
            else:
                current['results'].append({'subitem': line.strip(), 'result': ''})
        else:
            current = None
    return sorted(out, key=lambda t: t['done'], reverse=True)


def discarded_tasks(xlsx=XLSX_PATH):
    """The sheet's Discarded tasks (set in Excel; tasks.md no longer has them), shaped
    like finished_tasks' with "discarded": True."""
    out = {}
    for r in read_sheet(xlsx):
        if r.get('Status') != 'Discarded' or not r.get('Items'):
            continue
        t = out.setdefault(r['Items'], {'text': r['Items'], 'added': _day(r.get('Begin')), 'due': _day(r.get('Due')),
                                        'done': '', 'results': [], 'discarded': True})
        if r.get('Subitem') or r.get('Result'):
            t['results'].append({'subitem': str(r.get('Subitem') or ''), 'result': str(r.get('Result') or '')})
    return list(out.values())


def _day(value):
    value = as_date(value)
    return value.strftime('%Y-%m-%d') if isinstance(value, datetime) else ''


def _write_lines(path, lines, bom):
    with open(path, 'wb') as f:
        f.write((b'\xef\xbb\xbf' if bom else b'') + ''.join(lines).encode('utf-8'))


def restore_task(path, title):
    """A finished task back to open: `- [x]` to `- [ ]`, without the result lines its
    finishing added (those dated its finish day, the newest date among them); results
    from before stay. The sync then makes it Pending."""
    with open(path, 'rb') as f:
        raw = f.read()
    bom = raw.startswith(b'\xef\xbb\xbf')
    lines = raw.decode('utf-8-sig').splitlines(keepends=True)
    start = next((i for i, line in enumerate(lines)
                  if (m := re.match(r'^- \[x\] (.+)$', line.rstrip('\r\n'), re.I)) and strip_tags(m.group(1)) == title), None)
    if start is None:
        raise ValueError('finished task not found in tasks.md')
    end = start + 1
    while end < len(lines) and re.match(r'^\s+\S', lines[end]):
        end += 1
    day = lambda line: (re.search(r'\((\d{4}-\d{2}-\d{2})\)', line) or [None, ''])[1]
    results = [i for i in range(start + 1, end) if RESULT_VALUE.match(lines[i])]
    done = max((day(lines[i]) for i in results), default='')
    for i in reversed(results):
        if done and day(lines[i]) == done:
            del lines[i]
    lines[start] = re.sub(r'^- \[x\]', '- [ ]', lines[start], flags=re.I)
    _write_lines(path, lines, bom)


def restore_discarded(path, xlsx, title):
    """A Discarded task back to open: written at the end of tasks.md with its begin and
    due dates, sub-items and results, and its Discarded rows taken out of the sheet (the
    sync would keep them as well). The sync then makes it Pending."""
    t = next((t for t in discarded_tasks(xlsx) if t['text'] == title), None)
    if not t:
        raise ValueError('discarded task not found in TaskDatabase.xlsx')
    new = [f"- [ ] {title}" + (f" (added:{t['added']})" if t['added'] else '') +
           (f" [due:{t['due']}]" if t['due'] else '') + '\n']
    for r in t['results']:
        if r['subitem']:
            new.append(f"      {r['subitem']}\n")
        if r['result']:
            new.append(f"      → 結果：{r['result']}\n")
    with open(path, 'rb') as f:
        raw = f.read()
    bom = raw.startswith(b'\xef\xbb\xbf')
    lines = raw.decode('utf-8-sig').splitlines(keepends=True)
    if lines and not lines[-1].endswith('\n'):
        lines[-1] += '\n'
    _write_lines(path, lines + new, bom)
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]
    cols = {str(c.value): c.column for c in ws[1]}
    for r in reversed(range(2, ws.max_row + 1)):
        if ws.cell(r, cols['Items']).value == title and ws.cell(r, cols['Status']).value == 'Discarded':
            ws.delete_rows(r)
    wb.save(xlsx)


def read_sheet(path=XLSX_PATH):
    """Rows of the first worksheet as dicts keyed by the header row, or [] when unreadable."""
    if not os.path.exists(path):
        return []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
    except Exception:
        return []
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = ['' if h is None else str(h) for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:]]


def status_map(path=XLSX_PATH):
    out = {}
    for row in read_sheet(path):
        items = row.get('Items')
        if items is not None and items not in out:
            out[items] = row.get('Status')
    return out


def parse_date(text):
    try:
        return datetime.strptime(text, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None


def priority_score(due, status, today):
    if status == 'Finished':
        return 999999
    weight = {'Checking': 1, 'Ongoing': 2, 'Pending': 3}.get(status, 3)
    if due is None:
        return 999998
    diff = (due - today).days
    sign = -1 if due < today else 1
    return sign * abs(diff) * weight


def display_tasks(path=TASKS_PATH, count=5, today=None):
    """(top tasks in priority order, how many more there are)."""
    today = today or date.today()
    tasks = unchecked_tasks(path)
    statuses = status_map(os.path.join(os.path.dirname(path), 'TaskDatabase.xlsx'))
    for t in tasks:
        t['status'] = statuses.get(t['text'])
        due = parse_date(t['due']) if t['due'] else None
        if due is None and not t['due'] and t['added']:
            added = parse_date(t['added'])
            due = added + timedelta(days=7) if added else None
        t['score'] = priority_score(due, t['status'], today)
    ranked = sorted(tasks, key=lambda t: t['score'])
    top = ranked[:count]
    return top, len(ranked) - len(top)


def hud_message(path=TASKS_PATH, count=5, today=None):
    top, remaining = display_tasks(path, count, today)
    if not top:
        return '目前沒有待辦事項'
    lines = []
    for n, t in enumerate(top, 1):
        title = f"{n}. {t['text']}"
        if t['due']:
            title += f" ({t['due']})"
        if t['status']:
            title += f" [{t['status']}]"
        lines.append(title)
        lines += [f'      • {d}' for d in t['description']]
    message = '\n'.join(lines)
    if remaining > 0:
        message += f'\n\n還有 {remaining} 項，開 Claude 問我'
    return message


# --- edits made from Agent Mission Control ------------------------------------
# A task's due date lives in tasks.md ([due:...]) and its status in the Status
# column of TaskDatabase.xlsx (matched on Items = the task title). Finishing goes
# through tasks.md - tick the task and let the sync set Finished - as a chat edit does.

def edit_task_line(path, title, change):
    with open(path, 'rb') as f:
        raw = f.read()
    bom = raw.startswith(b'\xef\xbb\xbf')
    lines = raw.decode('utf-8-sig').splitlines(keepends=True)
    for i, line in enumerate(lines):
        body = line.rstrip('\r\n')
        m = re.match(r'^- \[ \] (.+)$', body)
        if m and strip_tags(m.group(1)) == title:
            lines[i] = change(body) + line[len(body):]
            break
    else:
        raise ValueError('task not found in tasks.md')
    with open(path, 'wb') as f:
        f.write((b'\xef\xbb\xbf' if bom else b'') + ''.join(lines).encode('utf-8'))


def set_due(path, title, due):
    def change(b):
        if re.search(r'\[due:\S+\]', b):
            return re.sub(r'\[due:\S+\]', f'[due:{due}]', b)
        return f'{b} [due:{due}]'
    edit_task_line(path, title, change)


def add_result(path, title, text):
    """Before a task is ticked Finished: a `→ 結果 (YYYY-MM-DD)：<text>` line under each of
    its description lines that has no result yet, or, with no description lines, under
    the title (unless it has a result there already). A result already written stays."""
    with open(path, 'rb') as f:
        raw = f.read()
    bom = raw.startswith(b'\xef\xbb\xbf')
    lines = raw.decode('utf-8-sig').splitlines(keepends=True)
    start = next((i for i, line in enumerate(lines)
                  if (m := re.match(r'^- \[ \] (.+)$', line.rstrip('\r\n'))) and strip_tags(m.group(1)) == title), None)
    if start is None:
        raise ValueError('task not found in tasks.md')
    end = start + 1
    while end < len(lines) and re.match(r'^\s+\S', lines[end]):
        end += 1
    eol = lines[start][len(lines[start].rstrip('\r\n')):] or '\n'
    text = ' '.join(text.split()) or 'finished'
    result = lambda indent: f"{indent}→ 結果 ({date.today():%Y-%m-%d})：{text}{eol}"
    items = [i for i in range(start + 1, end) if not RESULT_LINE.match(lines[i])]
    if items:
        # From the last up, so the indexes before stay right.
        for i in reversed(items):
            if not (i + 1 < end and RESULT_LINE.match(lines[i + 1])):
                lines.insert(i + 1, result(re.match(r'^\s*', lines[i]).group(0)))
    elif end == start + 1:
        lines.insert(start + 1, result('      '))
    with open(path, 'wb') as f:
        f.write((b'\xef\xbb\xbf' if bom else b'') + ''.join(lines).encode('utf-8'))


def discard_task(path, xlsx, title):
    """Discards the open task `title`: its rows in the sheet become Discarded (the sync
    keeps those), then its line and the indented lines under it leave tasks.md. The
    task History lists it, and can restore it."""
    set_status(xlsx, title, 'Discarded')
    with open(path, 'rb') as f:
        raw = f.read()
    bom = raw.startswith(b'\xef\xbb\xbf')
    lines = raw.decode('utf-8-sig').splitlines(keepends=True)
    start = next((i for i, line in enumerate(lines)
                  if (m := re.match(r'^- \[ \] (.+)$', line.rstrip('\r\n'))) and strip_tags(m.group(1)) == title), None)
    if start is None:
        raise ValueError('task not found in tasks.md')
    end = start + 1
    while end < len(lines) and re.match(r'^\s+\S', lines[end]):
        end += 1
    del lines[start:end]
    _write_lines(path, lines, bom)


def set_checked(path, title):
    edit_task_line(path, title, lambda b: re.sub(r'^- \[ \]', '- [x]', b))


def write_status(xlsx, title, status):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]
    cols = {str(c.value): c.column for c in ws[1]}
    hits = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, cols['Items']).value == title:
            ws.cell(r, cols['Status']).value = status
            hits += 1
    if hits:
        wb.save(xlsx)
    return hits


def set_status(xlsx, title, status):
    if write_status(xlsx, title, status) == 0:
        # A fresh task is not in the sheet until the sync has run: run it, then write.
        sync()
        if write_status(xlsx, title, status) == 0:
            raise ValueError('task not found in TaskDatabase.xlsx')


# --- TaskDatabase.xlsx sync ---------------------------------------------------

def all_tasks(path):
    tasks, current = [], None
    for line in read_lines(path):
        m = re.match(r'^- \[( |x)\] (.+)$', line, re.I)
        result = RESULT_VALUE.match(line)
        if m:
            if current:
                tasks.append(current)
            body = m.group(2)
            added = re.search(r'\(added:(\S+)\)', body)
            due = re.search(r'\[due:(\S+)\]', body)
            current = {'text': strip_tags(body), 'added': added.group(1) if added else None,
                       'due': due.group(1) if due else None, 'checked': m.group(1).lower() == 'x',
                       'items': [], 'result': ''}
        elif current and result and current['items']:
            current['items'][-1]['result'] = result.group(1).strip()
        elif current and result:
            current['result'] = result.group(1).strip()
        elif current and re.match(r'^\s+\S', line):
            current['items'].append({'subitem': line.strip(), 'result': ''})
        else:
            if current:
                tasks.append(current)
            current = None
    if current:
        tasks.append(current)
    return tasks


def as_date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    return value


def sync(tasks_path=TASKS_PATH, xlsx_path=XLSX_PATH):
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    existing = read_sheet(xlsx_path)
    statuses = {f"{r.get('Items') or ''}|{r.get('Subitem') or ''}": r.get('Status') for r in existing}
    discarded = [r for r in existing if r.get('Status') == 'Discarded']

    rows = []
    for t in all_tasks(tasks_path):
        items = t['items'] or [{'subitem': '', 'result': t['result']}]
        for item in items:
            if t['checked']:
                status = 'Finished'
            else:
                prior = statuses.get(f"{t['text']}|{item['subitem']}")
                status = prior if prior and prior != 'Finished' else 'Pending'
            begin = datetime.strptime(t['added'], '%Y-%m-%d') if t['added'] else None
            if t['due']:
                due = datetime.strptime(t['due'], '%Y-%m-%d')
            else:
                due = begin + timedelta(days=7) if begin else None
            rows.append([None, begin, due, status, t['text'], item['subitem'], item['result']])
    for r in discarded:
        rows.append([None, as_date(r.get('Begin')), as_date(r.get('Due')), 'Discarded',
                     r.get('Items'), r.get('Subitem'), r.get('Result')])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tasks'
    header = ['P', 'Begin', 'Due', 'Status', 'Items', 'Subitem', 'Result']
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    for i, row in enumerate(rows):
        n = i + 2
        row[0] = (f'=IF(OR(D{n}="Finished",D{n}="Discarded"),999999,IF(C{n}<TODAY(),-1,1)'
                  f'*ABS(C{n}-TODAY())*_xlfn.IFS(D{n}="Checking",1,D{n}="Ongoing",2,D{n}="Pending",3))')
        ws.append(row)
        ws.cell(n, 2).number_format = 'yyyy-mm-dd'
        ws.cell(n, 3).number_format = 'yyyy-mm-dd'
    for c in range(1, len(header) + 1):
        width = max(len(str(ws.cell(r, c).value or '')) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(c)].width = 12 if c in (2, 3) else min(max(width + 2, 6), 80)
    tmp = xlsx_path + '.tmp'
    wb.save(tmp)
    os.replace(tmp, xlsx_path)
    return len(rows)


# --- standalone popup -----------------------------------------------------------

def popup():
    import tkinter as tk
    sys.path.insert(0, os.path.join(os.path.dirname(HERE), 'Session HUD'))
    from single_instance import replace_previous
    replace_previous('todo-list-popup')

    bg, accent, title_c, desc_c, note_c = '#060a06', '#33ff66', '#8cff8c', '#46a05a', '#5a9664'
    root = tk.Tk()
    root.title('Todo List')
    root.geometry('680x320')
    root.configure(bg=bg)
    root.attributes('-topmost', True)
    try:
        root.attributes('-alpha', 0.92)
    except tk.TclError:
        pass
    tk.Frame(root, bg=accent, height=2).pack(fill='x')
    box = tk.Text(root, bg=bg, fg=title_c, font=('DejaVu Sans Mono', 12), relief='flat',
                  highlightthickness=0, wrap='word', padx=10, pady=8)
    box.pack(fill='both', expand=True)
    box.tag_configure('desc', foreground=desc_c)
    box.tag_configure('note', foreground=note_c, font=('DejaVu Sans Mono', 9, 'italic'))
    top, remaining = display_tasks()
    if not top:
        box.insert('end', '目前沒有待辦事項', 'note')
    for n, t in enumerate(top, 1):
        title = f"{n}. {t['text']}" + (f"  ({t['due']})" if t['due'] else '') + (f"  [{t['status']}]" if t['status'] else '')
        box.insert('end', ('' if n == 1 else '\n') + title)
        for d in t['description']:
            box.insert('end', f'\n      • {d}', 'desc')
    if remaining > 0:
        box.insert('end', f'\n\n還有 {remaining} 項，開 Claude 問我', 'note')
    box.configure(state='disabled')
    root.mainloop()


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'list'
    if cmd == 'sync':
        try:
            print(f'Synced {sync()} rows to {XLSX_PATH}')
        except Exception as e:
            print(f'同步失敗：{e}')
            sys.exit(1)
    elif cmd == 'popup':
        popup()
    else:
        print(hud_message())
